"""Regenerate Task 3 report EXACTLY in Task 1 style (Abhishek Jadhav)."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "report", "Internship_Task_3_Report.docx")

C_DARK = RGBColor(0x1F, 0x4E, 0x79)
C_BLUE = RGBColor(0x2E, 0x74, 0xB5)
C_GREY = RGBColor(0x59, 0x56, 0x59)

doc = Document()

# --- base styles to match Task 1 ---
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

h1 = doc.styles["Heading 1"]
h1.font.name = "Calibri"
h1.font.size = Pt(16)
h1.font.bold = True
h1.font.color.rgb = C_DARK

h2 = doc.styles["Heading 2"]
h2.font.name = "Calibri"
h2.font.size = Pt(12.5)
h2.font.bold = True
h2.font.color.rgb = C_BLUE

sec = doc.sections[0]
sec.top_margin = Inches(0.9)
sec.bottom_margin = Inches(0.9)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)

# --- footer: PAGE + label, 9pt grey centered (exact Task 1 pattern) ---
FOOT_LABEL = "Internship Task 3 \u2013 Placement Management System"
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
def ftrun(text="", bold=False, italic=False):
    r = fp.add_run(text)
    r.font.size = Pt(9)
    r.font.color.rgb = C_GREY
    r.bold = bold
    r.italic = italic
    return r
r = ftrun()
fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin"); r._r.append(fld1)
r = ftrun()
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"; r._r.append(instr)
r = ftrun()
fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "separate"); r._r.append(fld2)
r = ftrun("1")
r = ftrun()
fld3 = OxmlElement("w:fldChar"); fld3.set(qn("w:fldCharType"), "end"); r._r.append(fld3)
ftrun(f"  |  {FOOT_LABEL}")

# --- helpers ---
def centered(text, size, color=None, bold=False, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p

def cover_field(label, value):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(label)
    r1.font.size = Pt(12); r1.bold = True
    r2 = p.add_run(value)
    r2.font.size = Pt(12)
    return p

def h(t, lvl=1):
    doc.add_heading(t, level=lvl)

def p(t):
    return doc.add_paragraph(t)

def bullets(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")

def dtable(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, hh in enumerate(headers):
        t.rows[0].cells[j].text = hh
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = v
    doc.add_paragraph("")
    return t

def shotbox(title, desc):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2")
    tcPr.append(shd)
    p1 = cell.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(title); r1.bold = True; r1.font.color.rgb = C_DARK
    p2 = cell.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(desc); r2.italic = True; r2.font.size = Pt(10)
    doc.add_paragraph("")
    return t

# ================= COVER (exact Task 1 layout) =================
doc.add_paragraph("")
doc.add_paragraph("")
doc.add_paragraph("")
centered("INTERNSHIP TASK 3", 14, C_BLUE, bold=True)
centered("Internship & Placement System", 28, C_DARK, bold=True)
centered("Student Internship & Placement Management System \u2013 Practical Implementation Project", 13, C_GREY)
doc.add_paragraph("")
cover_field("Name: ", "Abhishek Jadhav")
cover_field("Role: ", "Full Stack Developer Intern")
cover_field("Project: ", "Student Internship & Placement Management System")
cover_field("Date: ", "09 October 2026")
doc.add_paragraph("")
centered("Web Development Internship \u00b7 Flask, SQLite, HTML5, CSS3, Vanilla JavaScript", 10, C_GREY, italic=True)
doc.add_page_break()

# ================= TOC =================
h("Table of Contents", 1)
for item in ["1. Cover Page", "2. Abstract", "3. Introduction and Problem Statement",
"4. Project Objectives and Scope", "5. Functional and Non-Functional Requirements",
"6. Technology Stack and Justification", "7. System Architecture", "8. Database Schema and Relationships",
"9. Implementation Details and Modules", "10. Authentication and Security",
"11. User Interface Screenshots", "12. Testing Methodology and Test Cases",
"13. Actual Test Results", "14. Bugs Discovered and Corrections",
"15. Performance and Reliability Considerations", "16. Challenges and Learning Outcomes",
"17. Installation and Execution Guide", "18. Limitations and Future Improvements",
"19. Conclusion", "20. References"]:
    doc.add_paragraph(item)
pn = doc.add_paragraph()
r0 = pn.add_run("Note: ")
r0.font.size = Pt(10)
r1 = pn.add_run("Update page numbers in Microsoft Word via Insert > Page Number if needed. Page numbers are already added in the footer.")
r1.italic = True
r1.font.size = Pt(10)
doc.add_paragraph("")

DESC_LONG = (
"This Student Internship and Placement Management System is a full-stack web application designed to help students discover internship and entry-level placement opportunities, submit structured applications, track each application's status over time, and steadily improve their employability through skills and resume management, while giving administrators a controlled back-office to publish listings, review applicants, and export operational reports. Students register and log in to a personal dashboard that summarises total applications, applications under review, shortlisted and selected counts, upcoming application deadlines for opportunities they have not yet applied to, and recent status activity drawn from an auditable history table, so progress is visible at a glance without exposing any other student's private data. Opportunity discovery supports keyword search across title, company and description plus filters for location, skill, internship type, remote, on-site and hybrid mode, paid versus unpaid, and deadline ordering, with each listing showing eligibility, required skills, duration, stipend, and application instructions, and with all sample listings explicitly labelled as training data rather than verified live vacancies. Applications are created through an internal form requiring a meaningful cover letter, are blocked at both the application and database level from duplicating the same student-opportunity pair, and move through the Applied, Under Review, Shortlisted, Selected and Rejected states with timestamps and free-text reviewer notes recorded in a separate history table for transparency. Profiles combine contact and education fields with a many-to-many technical-skills inventory and a printable resume view plus an optional strictly validated file upload, and administrators manage users, listings, application reviews, aggregate statistics and filtered CSV exports through server-validated routes protected by role checks. The system is intentionally built on a lightweight Flask and SQLite stack with server-rendered templates, parameterised queries, session-based authentication with per-session CSRF tokens, and an automated pytest suite covering registration, permissions, search, duplicates, status transitions, cross-student access control, CSV export, invalid input, empty states and date validation, making it suitable as a maintainable, secure-by-default internship project that can be installed locally with pip, seeded idempotently, and extended toward notifications, analytics and production-grade deployment."
)

h("2. Abstract", 1)
p("This report documents the design, implementation, testing and documentation of Task 3: a Student Internship & Placement Management System. It is a separate, original project from Task 2 (Student Attendance Management System in student-attendance-system/, which was inspected and left untouched). The system provides student registration and role-based login, searchable internship listings with an admin publishing workflow, duplicate-safe applications with auditable status transitions, a personal dashboard, skills and printable-resume management with strictly validated uploads, an admin back-office, and filtered CSV reporting. It was built with Flask and SQLite, verified by 13 automated pytest cases (all passing), seeded idempotently, and prepared for GitHub without secrets. Sample listings are synthetic training data. No screenshots, features or test outcomes are fabricated.")

h("3. Introduction and Problem Statement", 1)
p("Task 3, titled Practical Implementation Project, is the advanced full-stack milestone of the Web Development internship. After Task 1 (Fundamentals and Setup) and Task 2 (Attendance Management), it requires requirements analysis, architecture, frontend-backend integration, authentication, CRUD, testing, security, performance and professional documentation.")
p("Students typically discover internships through scattered posts, spreadsheets and chat forwards, losing track of deadlines and feedback, while placement cells lack a single view of who applied where and with what outcome. This project therefore provides a central portal where students find opportunities, apply uniformly, monitor status and prepare profiles, while admins curate listings and review pipelines with timestamps and exports. Trust matters as much as convenience: students must only see their own data, and uploads must be safe.")
p("The web project (folder: student-internship-placement-system) and this Word report are maintained as separate deliverables, in the same way Task 1 kept internship-task-1 apart from its report.")

h("4. Project Objectives and Scope", 1)
p("The official milestone for this task was to build a functional full-stack placement application demonstrating planning, architecture, integration, auth, CRUD, testing, security, performance and documentation. To meet it, the work focused on the following areas:")
bullets(["providing secure student and admin accounts with session authentication and CSRF;",
"publishing searchable, filterable, deadline-aware listings with admin CRUD and validation;",
"accepting internal applications, blocking duplicates, and tracking five statuses with history;",
"showing a personal dashboard (counts, deadlines, recent activity) with strict ownership;",
"managing skills and education, a printable resume, and an optional safe file upload;",
"offering admin user, listing, review and statistics workflows plus filtered CSV reports;",
"verifying everything with pytest and documenting only verified behaviour."])
p("Out of scope: real job-board scraping, email or SMS notifications, payments and production hosting. These are recorded as future work.")

h("5. Functional and Non-Functional Requirements", 1)
dtable(["ID", "Requirement", "Type"], [
["F1", "Register, login, logout; first user bootstraps admin", "Functional"],
["F2", "Search title/company/skill/location/type; filter mode/paid/deadline", "Functional"],
["F3", "Admin create/edit/publish/archive/delete with date validation", "Functional"],
["F4", "Apply with cover letter; UNIQUE(user, opportunity) blocks duplicates", "Functional"],
["F5", "Statuses Applied / Under Review / Shortlisted / Selected / Rejected + history", "Functional"],
["F6", "Student dashboard counts + deadlines + recent activity (own data only)", "Functional"],
["F7", "Skills CRUD, education profile, printable resume, validated upload", "Functional"],
["F8", "Admin users/listings/reviews/stats; reports filter + CSV export", "Functional"],
["N1", "Hashed passwords, parameterised SQL, CSRF, RBAC, IDOR protection", "Non-functional (security)"],
["N2", "Responsive, accessible forms, empty states, flash messages; every button works", "Non-functional (usability)"],
["N3", "pytest coverage of core flows; idempotent seed; indexed queries", "Non-functional (reliability)"]])

h("6. Technology Stack and Justification", 1)
p("Only the tools actually used for this task are listed below. No versions are invented.")
dtable(["Tool / Technology", "Purpose in Task 3", "Details"], [
["Python 3.14", "Backend runtime", "Runs Flask app, seed and report scripts"],
["Flask 3.1", "Web framework", "Routing, sessions, Jinja templates"],
["SQLite", "Relational database", "File placement.db; FKs, UNIQUE, indexes"],
["HTML5", "Page structure", "15 Jinja templates extending base.html"],
["CSS3", "Page styling", "static/style.css, no framework; responsive grid"],
["JavaScript (Vanilla)", "Minor UI behaviour", "static/app.js delete confirmations"],
["pytest 9.1", "Automated testing", "13 tests in tests/test_app.py"],
["python-docx 1.2", "Report generation", "Generates this .docx file"],
["Git / GitHub", "Version control", "Repo init, commit, push (URL pasted in \u00a720)"]])
p("Reuse of installed Flask, pytest and python-docx avoided new installs. No Docker, WSL, paid APIs or heavyweight frameworks were used, as required.")

h("7. System Architecture", 1)
p("Monolithic Flask app with clearly separated route sections (auth, opportunities, applications, profiles, administration, reporting), plus config.py, schema.sql and seed.py. Server-rendered Jinja templates share base.html; one CSS file handles layout; a tiny JS file adds confirmations. Request flow: route \u2192 login_required / admin_required \u2192 server validation \u2192 parameterised SQL in a transaction \u2192 redirect + flash \u2192 render. Sessions store user_id, role and _csrf; cookies are HttpOnly + SameSite=Lax. Full design is also recorded in ARCHITECTURE.md.")
h("7.1 Module Map", 2)
dtable(["File", "Purpose"], [
["app.py", "create_app factory, get_db/init_db, CSRF, decorators, all routes, error pages"],
["config.py", "SECRET_KEY, DATABASE, UPLOAD_FOLDER, extension whitelist, size cap"],
["schema.sql", "DDL for 7 tables with keys, constraints and indexes"],
["seed.py", "Idempotent demo data (admin, student, skills, 5 sample listings)"],
["templates/ (15)", "Landing, auth, dashboard, discovery, detail, apply, tracking, profile, resume, admin, reports"],
["static/", "style.css responsive theme; app.js confirmations"]])

h("8. Database Schema and Relationships", 1)
p("Relational SQLite design with primary keys, foreign keys, unique constraints, indexes and transactions where appropriate.")
dtable(["Table", "Key columns / constraints"], [
["users", "id PK; email UNIQUE NOCASE; password_hash; role check; idx email/role"],
["student_profiles", "user_id UNIQUE FK CASCADE; education/contact/resume fields"],
["skills", "id PK; name UNIQUE NOCASE"],
["student_skills", "user FK; skill FK; proficiency check; UNIQUE(user, skill)"],
["opportunities", "title/company/location/type/mode/paid/stipend/duration/skills/eligibility/description/instructions/deadline/status check; idx status/deadline/company/mode"],
["applications", "user FK; opportunity FK; cover_letter; status check; timestamps; UNIQUE(user, opp)"],
["application_status_history", "application FK CASCADE; old/new status; changed_by; note; changed_at"]])
h("8.1 Relationships", 2)
p("User 1\u20141 profile; user N\u2014M skills via student_skills; user N\u2014M opportunities via applications; application 1\u2014N history. PRAGMA foreign_keys=ON per connection. UNIQUE(user, opportunity) enforces the no-duplicate rule at the database layer, and the history table (not just a column) provides the audit trail behind the dashboard activity feed.")

h("9. Implementation Details and Modules", 1)
h("9.1 Authentication and Profiles", 2)
p("Register validates name, email regex, password length and confirmation, then rejects duplicate emails (400 with flashes); the very first user becomes admin so a fresh database is usable. Login verifies the Werkzeug hash, clears and regenerates the session plus CSRF token, and honours a safe relative next parameter. Profiles auto-create a row, update only whitelisted fields with length caps, add skills via INSERT OR IGNORE plus link (duplicate \u2192 flash), and print from resume.html with print CSS.")
h("9.2 Opportunities and Applications", 2)
p("Students query published listings only; admins may add a status filter. Keyword LIKE covers title/company/description; equality covers type/mode; paid toggles; skill LIKE covers skills_required; ordering is deadline ASC with applied-set badges. Detail hides non-published rows from students (404). Apply uses GET form plus POST requiring a 10+ character cover letter, rejects past deadlines, and catches IntegrityError as 'already applied'. Status POSTs rewrite the row and append history with reviewer notes and timestamps.")
h("9.3 Administration and Reporting", 2)
p("Admin home shows counts, GROUP BY status and the 8 most recent applications. Users list shows per-student application counts with delete (never self). Listing new/edit share one validator (title, company, enums, YYYY-MM-DD deadline). Review filters by status; reports add from/to date filters auto-scoped to the student for non-admins; CSV uses the csv module with a download header. Invalid filters flash and redirect; empty results render honest empty states.")

h("10. Authentication and Security", 1)
bullets(["password hashing with Werkzeug; HttpOnly + SameSite=Lax session cookies; secret from environment;",
"per-session CSRF token validated on every POST (400 on mismatch; covered by test_csrf_enforced);",
"login_required and admin_required decorators plus per-record ownership checks (non-owner /applications/<id> \u2192 403);",
"parameterised SQL throughout; no string-interpolated queries;",
"server-side validation (email regex, lengths, enums, valid_date) and safe 400/403/404 pages;",
"resume uploads restricted to .pdf/.doc/.docx/.txt, secure_filename plus random prefix, 2 MB cap, stored outside templates and never executed."])
p("Verified by test_idor_blocked: changing the URL to another student's application id returns 403 when logged in as a different student.")

h("11. User Interface Screenshots", 1)
p("Pages delivered: landing (hero + counts + featured), register, login, student dashboard (4 stat cards + deadlines + activity), discovery (filters + cards + applied badges), detail, apply form, my applications, application detail with history, profile (education + skills + upload), printable resume, admin dashboard, users, listing form, review table, reports with CSV. Layout is responsive with consistent navigation, empty states, validation errors and confirmation flashes. Sample listings carry an explicit training-data notice.")
shotbox("[INSERT SCREENSHOT 1 \u2013 STUDENT DASHBOARD]", "Show dashboard stat cards, upcoming deadlines and recent activity for student@example.com.")
shotbox("[INSERT SCREENSHOT 2 \u2013 OPPORTUNITY DISCOVERY]", "Show search filters and result cards with applied badges.")
shotbox("[INSERT SCREENSHOT 3 \u2013 APPLICATION DETAIL + PYTEST]", "Show status history; alongside terminal output of pytest 13 passed.")
shotbox("[INSERT SCREENSHOT 4 \u2013 PROFILE + PRINTABLE RESUME]", "Show skills editor and resume print preview.")
shotbox("[INSERT SCREENSHOT 5 \u2013 ADMIN + REPORTS CSV]", "Show admin stats, review table and exported applications.csv.")
p("Note: placeholders follow the Task 1 report convention. Paste genuine captures before submission; no fabricated screenshots are included.")

h("12. Testing Methodology and Test Cases", 1)
p("pytest with a temporary SQLite database per test and CSRF tokens scraped from GET forms (enforcement stays ON). Each Task \u00a77 bullet maps to at least one test:")
dtable(["#", "Test", "Expected Result", "Status"], [
["1", "test_register_and_login", "Admin + student register; student dashboard 200", "PASS"],
["2", "test_invalid_credentials", "Wrong password \u2192 401", "PASS"],
["3", "test_invalid_register", "Bad email / short pw / mismatch \u2192 400", "PASS"],
["4", "test_admin_can_create_and_edit_opp", "Admin POST creates; appears in list", "PASS"],
["5", "test_student_cannot_create_opp", "Student GET /admin* \u2192 403", "PASS"],
["6", "test_search_filter", "q=Python hits; nonsense q shows empty state", "PASS"],
["7", "test_duplicate_application_blocked", "Second POST blocked by UNIQUE", "PASS"],
["8", "test_status_transition_and_history", "Admin Shortlisted; visible in detail", "PASS"],
["9", "test_idor_blocked", "Non-owner /applications/1 \u2192 403", "PASS"],
["10", "test_csv_export", "200 + text/csv header", "PASS"],
["11", "test_invalid_opp_date_rejected", "Bad deadline \u2192 400", "PASS"],
["12", "test_empty_db_states", "Empty lists render honest empty states", "PASS"],
["13", "test_csrf_enforced", "POST without token \u2192 400", "PASS"]])

h("13. Actual Test Results", 1)
p("Command: python -m pytest tests -v. Result: 13 passed in ~13s (Windows, Python 3.14.4, pytest 9.1.1). All 13 rows in the table above report PASS; the full names are test_register_and_login, test_invalid_credentials, test_invalid_register, test_admin_can_create_and_edit_opp, test_student_cannot_create_opp, test_search_filter, test_duplicate_application_blocked, test_status_transition_and_history, test_idor_blocked, test_csv_export, test_invalid_opp_date_rejected, test_empty_db_states and test_csrf_enforced. The seed script was run twice: the second run printed 'exists, skip' for every row, proving idempotence. No results are assumed.")

h("14. Bugs Discovered and Corrections", 1)
dtable(["#", "Symptom", "Cause \u2192 Correction"], [
["1", "test_register_and_login: /dashboard 302, not 200", "First user is admin by design \u2192 test now registers admin + student, asserts student dashboard"],
["2", "Status/IDOR tests returned 404", "Fixture titles ('T'/'C'/'S') failed min-length validation so rows were never created \u2192 valid fixtures"],
["3", "Status test 404 on /admin/applications/1", "Detail route is /applications/<id>; test used list URL \u2192 corrected test URL"],
["4", "Whole test file corrupted", "replaceAll 'T' rewrote every T character \u2192 rewrote file cleanly, re-ran green"],
["5", "Seed duplication concern", "Verified INSERT OR IGNORE + existence checks; double-run safe"]])

h("15. Performance and Reliability Considerations", 1)
p("By completing this task, the following reliability practices were applied:")
bullets(["indexes on users.email, opportunities(status/deadline/company/mode), applications(user/opp/status) and history(app) keep filters fast;",
"LIMIT 3/5/8 on featured listings, deadlines and activity avoids full scans;",
"single SELECT COUNT(*) queries and JOINed list pages avoid N+1 queries;",
"one transaction per write with UNIQUE constraints as last-line defence;",
"safe error pages without stack leaks; optimisations applied only where justified, with no premature caching."])

h("16. Challenges and Learning Outcomes", 1)
p("By completing this task, I learned:")
bullets(["how to structure a Flask app with a create_app factory and temp-database pytest fixtures;",
"how to enforce integrity at the database layer (UNIQUE pairs, FKs) instead of relying on UI checks;",
"how to implement per-session CSRF and per-record ownership checks that survive URL tampering;",
"how to keep validation, flashes and empty states consistent across 15 templates;",
"how to generate an academic Word report with python-docx and document only executed results;",
"how to scope a project (deferring notifications, ranking and hosting) without breaking core flows."])
p("The main challenge was balancing speed with correctness: most 'app bugs' turned out to be fixture bugs tripping the very validation under test, and one bulk string replacement corrupted a file \u2014 both fixed by slowing down, reading the error, and re-running the suite.")

h("17. Installation and Execution Guide", 1)
dtable(["Step", "Command"], [
["1", 'cd "student-internship-placement-system"'],
["2", "pip install -r requirements.txt"],
["3", "python seed.py"],
["4", "python app.py  \u2192  http://127.0.0.1:5000"],
["5", "python -m pytest tests -v"]])
p("Seed creates placement.db, admin@example.com / admin123, student@example.com / student123 and 5 sample listings; re-running is safe. Copy .env.example to .env to set a real SECRET_KEY on any shared host. Troubleshooting: 'no such table' \u2192 run seed.py; CSRF 400 \u2192 submit via rendered forms; port busy \u2192 change port in app.py; uploads land in uploads/ (git-ignored).")

h("18. Limitations and Future Improvements", 1)
p("Limitations: no email notifications, no full-text ranking, no pagination beyond LIMITs, no avatar handling, screenshots captured manually, single-process dev server. Future work: pagination and sorting UI, deadline reminders, saved listings, a recruiter role, analytics charts, magic-byte file sniffing, rate limiting, production WSGI with migrations, and CI running pytest on every push.")

h("19. Conclusion", 1)
p("Task 3 successfully delivers the advanced full-stack milestone: secure auth, searchable listings, duplicate-safe applications with history, personal and admin dashboards, skills and resume workflows, filtered CSV reports, 13/13 passing tests, an idempotent seed and honest documentation. Task 2 remains untouched. Keeping the scope tight ensured every button, form, filter and workflow genuinely works, and the project is ready for local demo and GitHub publication with clear paths for extension.")

h("20. References", 1)
bullets(["Flask documentation (flask.palletsprojects.com) \u2014 routing, sessions, testing.",
"SQLite documentation (sqlite.org) \u2014 foreign keys, UNIQUE, indexes.",
"Werkzeug \u2014 password hashing, secure_filename.",
"pytest documentation \u2014 fixtures, temporary databases.",
"python-docx documentation \u2014 Word generation.",
"Course Task 3 brief (\u00a7\u00a71\u201312) \u2014 scope and acceptance criteria."])

h("GitHub Repository", 1)
p("Provide your real repository URL after you create and push the project. Do not invent the URL.")
p("GitHub Repository: [PASTE YOUR ACTUAL GITHUB URL HERE]")
pn2 = doc.add_paragraph()
pn2.add_run("Note: ").font.size = Pt(10)
rN = pn2.add_run("Example format only (do not use this as your link): https://github.com/YOUR-USERNAME/internship-placement-system \u2013 replace with your real link.")
rN.italic = True
rN.font.size = Pt(10)

h("Appendix A \u2013 Internship-Portal Project Description (>200 words)", 1)
p(DESC_LONG)
pn3 = doc.add_paragraph()
rW = pn3.add_run(f"Word count of description: {len(DESC_LONG.split())} words.")
rW.italic = True

doc.core_properties.title = "Internship Task 3 \u2013 Placement Management System"
doc.core_properties.author = "Abhishek Jadhav"
doc.core_properties.subject = "Student Internship & Placement Management System"

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
body_words = len(" ".join([x.text for x in doc.paragraphs]).split())
print(f"Saved {OUT} | body words={body_words} | desc words={len(DESC_LONG.split())}")
